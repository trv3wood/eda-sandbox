from __future__ import annotations

import json
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any

from .io import (
    dump_json,
    file_digest,
    load_yaml,
    project_paths,
    relative_to_project,
    resolve_inputs,
)


@dataclass(frozen=True)
class Evidence:
    id: str
    source_kind: str
    source_path: str
    source_digest: str
    locator: str
    text: str
    extractor: str


def _evidence(
    *,
    kind: str,
    path: Path,
    project_dir: Path,
    locator: str,
    text: str,
    extractor: str,
) -> Evidence:
    import hashlib

    normalized = " ".join(text.split())
    source_path = relative_to_project(project_dir, path)
    source_digest = file_digest(path)
    key = (
        f"{kind}\0{source_path}\0{source_digest}\0{locator}\0{normalized}".encode()
    )
    evidence_id = "ev-" + hashlib.sha256(key).hexdigest()[:16]
    return Evidence(
        id=evidence_id,
        source_kind=kind,
        source_path=source_path,
        source_digest=source_digest,
        locator=locator,
        text=normalized,
        extractor=extractor,
    )


def extract_docx(path: Path, project_dir: Path) -> tuple[list[Evidence], dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX extraction requires python-docx") from exc

    document = Document(path)
    evidence: list[Evidence] = []
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        item = _evidence(
            kind="docx",
            path=path,
            project_dir=project_dir,
            locator=f"paragraph:{index}",
            text=text,
            extractor="python-docx",
        )
        evidence.append(item)
        paragraphs.append({"index": index, "style": paragraph.style.name, "evidence": item.id})

    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            values = []
            for column_index, cell in enumerate(row.cells, start=1):
                text = cell.text.strip()
                if not text:
                    values.append(None)
                    continue
                item = _evidence(
                    kind="docx",
                    path=path,
                    project_dir=project_dir,
                    locator=f"table:{table_index}/row:{row_index}/cell:{column_index}",
                    text=text,
                    extractor="python-docx",
                )
                evidence.append(item)
                values.append(item.id)
            rows.append(values)
        tables.append({"index": table_index, "rows": rows})
    return evidence, {"path": str(path), "paragraphs": paragraphs, "tables": tables}


def extract_markdown(path: Path, project_dir: Path) -> tuple[list[Evidence], dict[str, Any]]:
    """Extract non-empty Markdown lines as source-located specification evidence."""
    evidence: list[Evidence] = []
    lines = []
    for index, raw_line in enumerate(
        path.read_text(encoding="utf-8", errors="replace").splitlines(), start=1
    ):
        text = raw_line.strip()
        if not text:
            continue
        item = _evidence(
            kind="markdown",
            path=path,
            project_dir=project_dir,
            locator=f"line:{index}",
            text=text,
            extractor="systemc-tlm-agent-markdown",
        )
        evidence.append(item)
        lines.append({"line": index, "evidence": item.id})
    return evidence, {"path": str(path), "lines": lines}


def extract_document(path: Path, project_dir: Path) -> tuple[list[Evidence], dict[str, Any]]:
    """Dispatch supported specification documents by their file extension."""
    if path.suffix.lower() == ".docx":
        return extract_docx(path, project_dir)
    if path.suffix.lower() in {".md", ".markdown"}:
        return extract_markdown(path, project_dir)
    raise ValueError(f"unsupported document format: {path}")


def extract_xlsx(path: Path, project_dir: Path) -> tuple[list[Evidence], dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX extraction requires openpyxl") from exc

    formula_book = load_workbook(path, data_only=False, read_only=False)
    value_book = load_workbook(path, data_only=True, read_only=False)
    evidence: list[Evidence] = []
    sheets: dict[str, Any] = {}
    for sheet in formula_book.worksheets:
        value_sheet = value_book[sheet.title]
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cached = value_sheet[cell.coordinate].value
                text = str(cell.value)
                if text.startswith("="):
                    text = f"formula={text}; cached={cached!r}"
                item = _evidence(
                    kind="xlsx",
                    path=path,
                    project_dir=project_dir,
                    locator=f"sheet:{sheet.title}/cell:{cell.coordinate}",
                    text=text,
                    extractor="openpyxl",
                )
                evidence.append(item)
                cells.append(
                    {
                        "cell": cell.coordinate,
                        "evidence": item.id,
                        "formula": cell.value if str(cell.value).startswith("=") else None,
                        "cached": cached,
                    }
                )
        sheets[sheet.title] = {
            "cells": cells,
            "merged_ranges": [str(value) for value in sheet.merged_cells.ranges],
        }
    return evidence, {"path": str(path), "sheets": sheets}


def extract_project(project_dir: Path, *, run_tools: bool = True) -> dict[str, Any]:
    """Extract immutable, source-located facts from manifest inputs."""
    paths = project_paths(project_dir)
    manifest = load_yaml(paths["manifest"])
    all_evidence: list[Evidence] = []

    documents = []
    for path in resolve_inputs(project_dir, manifest.get("documents", [])):
        evidence, facts = extract_document(path, project_dir)
        all_evidence.extend(evidence)
        documents.append(facts)

    registers = []
    for path in resolve_inputs(project_dir, manifest.get("registers", [])):
        evidence, facts = extract_xlsx(path, project_dir)
        all_evidence.extend(evidence)
        registers.append(facts)

    rtl_files = resolve_inputs(
        project_dir,
        manifest.get("rtl", []),
        directory_suffixes={".v", ".sv"},
    )
    testbench_files = resolve_inputs(
        project_dir,
        manifest.get("testbench", []),
        directory_suffixes={".v", ".sv"},
    )
    compile_config = manifest.get("eda_compile", {})
    if compile_config and not isinstance(compile_config, dict):
        raise ValueError("manifest eda_compile must be a mapping")
    compile_sources = resolve_inputs(
        project_dir,
        compile_config.get("sources", manifest.get("rtl", [])),
        directory_suffixes={".v", ".sv"},
    )
    include_dirs = []
    for value in compile_config.get("include_dirs", []):
        path = Path(value)
        resolved = (path if path.is_absolute() else project_dir / path).resolve()
        if not resolved.is_dir():
            raise FileNotFoundError(f"EDA include directory does not exist: {value}")
        include_dirs.append(resolved)
    defines = compile_config.get("defines", [])
    if not isinstance(defines, list) or not all(
        isinstance(value, str) and value for value in defines
    ):
        raise ValueError("manifest eda_compile.defines must be a list of strings")
    paths["facts"].mkdir(parents=True, exist_ok=True)
    dump_json(paths["facts"] / "documents.json", {"documents": documents})
    dump_json(paths["facts"] / "registers.json", {"workbooks": registers})
    target_top = manifest.get("target_top", manifest.get("top"))
    reference_top = manifest.get("reference_top") or target_top
    if not target_top:
        raise ValueError("manifest requires target_top (or legacy top)")
    rtl_facts = {
        "schema_version": 2,
        "status": "pending" if rtl_files else "missing",
        "backend": None,
        "target_top": target_top,
        "reference_top": reference_top,
        "top": target_top,
        "files": [],
        "design_rtl": [str(path) for path in rtl_files],
        "testbench": [str(path) for path in testbench_files],
        "compile_sources": [str(path) for path in compile_sources],
        "tools": {
            tool: {
                "status": "skipped",
                "reason": (
                    "no RTL inputs were provided"
                    if not rtl_files
                    else "EDA producer execution is pending"
                ),
            }
            for tool in (
                "surelog",
                "uhdm_elab",
                "uhdm_lint",
                "uhdm_hier",
                "uhdm",
                "verilator",
                "yosys",
            )
        },
    }
    dump_json(paths["facts"] / "rtl.json", rtl_facts)
    with paths["evidence"].open("w", encoding="utf-8") as stream:
        for item in sorted(all_evidence, key=lambda value: value.id):
            stream.write(json.dumps(asdict(item), ensure_ascii=False, sort_keys=True) + "\n")

    summary = {
        "evidence_count": len(all_evidence),
        "document_count": len(documents),
        "register_workbook_count": len(registers),
        "rtl_available": bool(rtl_files),
        "rtl_file_count": len(rtl_files),
        "testbench_file_count": len(testbench_files),
        "rtl_module_count": 0,
        "rtl_status": rtl_facts["status"],
        "missing_inputs": [] if rtl_files else ["rtl"],
    }
    dump_json(paths["facts"] / "summary.json", summary)
    if run_tools and rtl_files:
        from .tool_producers import finalize_tools, produce_rtl, produce_uhdm

        produce_uhdm(project_dir)
        produce_rtl(project_dir)
        finalize_tools(project_dir)
        return load_json(paths["facts"] / "summary.json")
    return summary
