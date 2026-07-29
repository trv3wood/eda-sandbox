from __future__ import annotations

import json
import os
import re
import argparse
import sys
from pathlib import Path
from typing import Any

from .schema import (
    GRAPH_SCHEMA_VERSION,
    SPEC_ENTITY_TYPES,
    SPEC_RELATION_TYPES,
    canonical_digest,
    entity,
    relationship,
    stable_id,
    write_jsonl,
)
from ..io import dump_json, file_digest, load_json, load_yaml, relative_to_project

SPEC_PROMPT_VERSION = "systemc-tlm-spec-graph/1"


def _source_ref(
    path: Path,
    project: Path,
    locator: str,
    *,
    start: int | None = None,
    end: int | None = None,
) -> dict[str, Any]:
    value: dict[str, Any] = {
        "source_path": relative_to_project(project, path),
        "source_digest": file_digest(path),
        "locator": locator,
    }
    if start is not None:
        value["start"] = start
    if end is not None:
        value["end"] = end
    return value


def _text_unit(
    path: Path,
    project: Path,
    locator: str,
    text: str,
    section_path: list[str],
) -> dict[str, Any]:
    normalized = text.strip()
    source = _source_ref(path, project, locator)
    return {
        "id": stable_id(
            "txt",
            {
                "source": source,
                "section_path": section_path,
                "text": normalized,
            },
        ),
        "text": normalized,
        "section_path": section_path,
        **source,
    }


def _markdown_document(
    path: Path, project: Path
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    sections: list[dict[str, Any]] = []
    units: list[dict[str, Any]] = []
    headings: list[str] = []
    block: list[str] = []
    block_start = 1

    def flush(end_line: int) -> None:
        nonlocal block
        text = "\n".join(block).strip()
        if text:
            units.append(
                _text_unit(
                    path,
                    project,
                    f"lines:{block_start}-{end_line}",
                    text,
                    list(headings),
                )
            )
        block = []

    for number, raw in enumerate(lines, 1):
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", raw)
        if match:
            flush(number - 1)
            level = len(match.group(1))
            title = match.group(2).strip()
            headings[:] = headings[: level - 1]
            headings.append(title)
            sections.append(
                {
                    "title": title,
                    "level": level,
                    "locator": f"line:{number}",
                    "section_path": list(headings),
                }
            )
            units.append(
                _text_unit(
                    path,
                    project,
                    f"line:{number}",
                    title,
                    list(headings),
                )
            )
            block_start = number + 1
        elif not raw.strip():
            flush(number - 1)
            block_start = number + 1
        else:
            if not block:
                block_start = number
            block.append(raw)
    flush(len(lines))
    return {
        "source_path": relative_to_project(project, path),
        "source_digest": file_digest(path),
        "kind": "markdown",
        "sections": sections,
    }, units


def _docx_document(
    path: Path, project: Path
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX graph extraction requires python-docx") from exc

    document = Document(path)
    sections: list[dict[str, Any]] = []
    units: list[dict[str, Any]] = []
    headings: list[str] = []
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name or ""
        match = re.match(r"Heading\s+([1-6])$", style, flags=re.IGNORECASE)
        if match:
            level = int(match.group(1))
            headings[:] = headings[: level - 1]
            headings.append(text)
            sections.append(
                {
                    "title": text,
                    "level": level,
                    "locator": f"paragraph:{index}",
                    "section_path": list(headings),
                }
            )
        unit = _text_unit(
            path, project, f"paragraph:{index}", text, list(headings)
        )
        units.append(unit)
        paragraphs.append(
            {"index": index, "style": style, "text_unit_id": unit["id"]}
        )
    tables = []
    for table_index, table in enumerate(document.tables, 1):
        rows = []
        for row_index, row in enumerate(table.rows, 1):
            cells = []
            for column_index, cell in enumerate(row.cells, 1):
                text = cell.text.strip()
                if not text:
                    cells.append(None)
                    continue
                locator = (
                    f"table:{table_index}/row:{row_index}/cell:{column_index}"
                )
                unit = _text_unit(path, project, locator, text, list(headings))
                units.append(unit)
                cells.append(unit["id"])
            rows.append(cells)
        tables.append({"index": table_index, "rows": rows})
    return {
        "source_path": relative_to_project(project, path),
        "source_digest": file_digest(path),
        "kind": "docx",
        "sections": sections,
        "paragraphs": paragraphs,
        "tables": tables,
    }, units


def extract_document_graph(
    project: Path, documents: list[Path]
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    trees = []
    units = []
    for path in documents:
        if path.suffix.lower() == ".docx":
            tree, values = _docx_document(path, project)
        elif path.suffix.lower() in {".md", ".markdown"}:
            tree, values = _markdown_document(path, project)
        else:
            raise ValueError(f"unsupported graph document format: {path}")
        trees.append(tree)
        units.extend(values)
    return {
        "schema_version": GRAPH_SCHEMA_VERSION,
        "documents": trees,
    }, units


def extract_workbook_units(
    project: Path, workbooks: list[Path]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not workbooks:
        return [], []
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX graph extraction requires openpyxl") from exc

    trees = []
    units = []
    for path in workbooks:
        book = load_workbook(path, data_only=False, read_only=False)
        sheets = []
        for sheet in book.worksheets:
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    locator = f"sheet:{sheet.title}/cell:{cell.coordinate}"
                    unit = _text_unit(
                        path,
                        project,
                        locator,
                        str(cell.value),
                        [sheet.title],
                    )
                    units.append(unit)
                    cells.append(
                        {"coordinate": cell.coordinate, "text_unit_id": unit["id"]}
                    )
            sheets.append({"name": sheet.title, "cells": cells})
        trees.append(
            {
                "source_path": relative_to_project(project, path),
                "source_digest": file_digest(path),
                "sheets": sheets,
            }
        )
    return trees, units


def spec_response_schema() -> dict[str, Any]:
    source_span = {
        "type": "object",
        "additionalProperties": False,
        "required": ["text_unit_id", "start", "end"],
        "properties": {
            "text_unit_id": {"type": "string"},
            "start": {"type": "integer", "minimum": 0},
            "end": {"type": "integer", "minimum": 1},
        },
    }
    return {
        "type": "object",
        "additionalProperties": False,
        "required": ["entities", "relationships"],
        "properties": {
            "entities": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": [
                        "local_id", "type", "name", "description", "source_spans"
                    ],
                    "properties": {
                        "local_id": {"type": "string"},
                        "type": {"enum": sorted(SPEC_ENTITY_TYPES)},
                        "name": {"type": "string", "minLength": 1},
                        "description": {"type": "string"},
                        "properties": {"type": "object"},
                        "source_spans": {
                            "type": "array",
                            "minItems": 1,
                            "items": source_span,
                        },
                    },
                },
            },
            "relationships": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": [
                        "type", "source_local_id", "target_local_id",
                        "description", "source_spans"
                    ],
                    "properties": {
                        "type": {"enum": sorted(SPEC_RELATION_TYPES)},
                        "source_local_id": {"type": "string"},
                        "target_local_id": {"type": "string"},
                        "description": {"type": "string"},
                        "properties": {"type": "object"},
                        "source_spans": {
                            "type": "array",
                            "minItems": 1,
                            "items": source_span,
                        },
                    },
                },
            },
        },
    }


def normalize_spec_response(
    response: dict[str, Any],
    text_units: list[dict[str, Any]],
    *,
    model: str,
    response_digest: str,
    system_fingerprint: str | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    try:
        import jsonschema
    except ImportError as exc:
        raise RuntimeError("Spec graph validation requires jsonschema") from exc
    jsonschema.validate(response, spec_response_schema())
    units = {item["id"]: item for item in text_units}

    def refs(spans: list[dict[str, Any]]) -> list[dict[str, Any]]:
        result = []
        for span in spans:
            unit = units.get(span["text_unit_id"])
            if unit is None:
                raise ValueError(f"unknown text unit {span['text_unit_id']}")
            start, end = span["start"], span["end"]
            if not (0 <= start < end <= len(unit["text"])):
                raise ValueError(f"invalid source span for {unit['id']}")
            result.append(
                {
                    "text_unit_id": unit["id"],
                    "source_path": unit["source_path"],
                    "source_digest": unit["source_digest"],
                    "locator": unit["locator"],
                    "start": start,
                    "end": end,
                    "text": unit["text"][start:end],
                }
            )
        return result

    provenance = {
        "extractor": SPEC_PROMPT_VERSION,
        "provider": "openai-compatible",
        "model": model,
        "response_digest": response_digest,
        "system_fingerprint": system_fingerprint,
        "source_backed": True,
    }
    entities = []
    local_to_id = {}
    for raw in response["entities"]:
        source_refs = refs(raw["source_spans"])
        value = entity(
            entity_type=raw["type"],
            name=raw["name"],
            description=raw["description"],
            properties=raw.get("properties", {}),
            source_refs=source_refs,
            provenance=provenance,
            identity={
                "type": raw["type"],
                "name": raw["name"],
                "source_refs": source_refs,
            },
        )
        if raw["local_id"] in local_to_id:
            raise ValueError(f"duplicate spec local_id {raw['local_id']}")
        local_to_id[raw["local_id"]] = value["id"]
        entities.append(value)
    relationships = []
    for raw in response["relationships"]:
        try:
            source_id = local_to_id[raw["source_local_id"]]
            target_id = local_to_id[raw["target_local_id"]]
        except KeyError as exc:
            raise ValueError(f"unknown spec relationship endpoint {exc.args[0]}") from exc
        relationships.append(
            relationship(
                relation_type=raw["type"],
                source_id=source_id,
                target_id=target_id,
                description=raw["description"],
                properties=raw.get("properties", {}),
                source_refs=refs(raw["source_spans"]),
                provenance=provenance,
            )
        )
    return entities, relationships


def produce_spec_graph(project: Path) -> dict[str, Any]:
    graph = project / ".systemc-agent" / "graph"
    manifest = load_json(graph / "manifest.json")
    text_units = sorted([
        json.loads(line)
        for line in (graph / "text_units.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ], key=lambda item: (item.get("ordinal", 0), item["id"]))
    if not text_units:
        write_jsonl(graph / "spec_entities.jsonl", [])
        write_jsonl(graph / "spec_relationships.jsonl", [])
        return {"status": "skipped", "reason": "no document text units"}
    config = load_yaml(project / "manifest.yaml").get("graph", {}).get(
        "spec_extraction", {}
    )
    model = config.get("model")
    base_url = config.get("base_url") or os.environ.get(
        config.get("base_url_env") or "SYSTEMC_TLM_LLM_BASE_URL"
    )
    api_key = os.environ.get(
        config.get("api_key_env") or "SYSTEMC_TLM_LLM_API_KEY", "not-required"
    )
    if not model or not base_url:
        raise ValueError(
            "graph.spec_extraction requires model and base_url or base_url_env"
        )
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("Spec graph production requires the graph extra") from exc

    prompt_digest = canonical_digest(
        {"version": SPEC_PROMPT_VERSION, "schema": spec_response_schema()}
    )
    input_digest = canonical_digest(text_units)
    cache = project / ".systemc-agent" / "tools" / "spec-llm-cache"
    cache.mkdir(parents=True, exist_ok=True)
    maximum = int(config.get("batch_max_chars", 24000))
    if maximum < 1000:
        raise ValueError("graph.spec_extraction.batch_max_chars must be >= 1000")
    batches: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    current_size = 0
    for unit in text_units:
        size = len(unit["text"])
        if current and current_size + size > maximum:
            batches.append(current)
            current, current_size = [], 0
        current.append(unit)
        current_size += size
    if current:
        batches.append(current)

    client = None
    entities: list[dict[str, Any]] = []
    relationships: list[dict[str, Any]] = []
    response_digests = []
    fingerprints = []
    for batch in batches:
        cache_key = canonical_digest({
            "input": canonical_digest(batch),
            "model": model,
            "prompt": prompt_digest,
            "temperature": 0,
        })
        cache_path = cache / f"{cache_key}.json"
        if cache_path.is_file():
            cached = load_json(cache_path)
        else:
            if client is None:
                client = OpenAI(base_url=base_url, api_key=api_key)
            completion = client.chat.completions.create(
                model=model,
                temperature=0,
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "Extract only source-grounded hardware specification "
                            "entities and relationships. Every item must cite exact "
                            "character spans."
                        ),
                    },
                    {
                        "role": "user",
                        "content": json.dumps(batch, ensure_ascii=False),
                    },
                ],
                response_format={
                    "type": "json_schema",
                    "json_schema": {
                        "name": "hardware_spec_graph",
                        "strict": True,
                        "schema": spec_response_schema(),
                    },
                },
            )
            content = completion.choices[0].message.content or "{}"
            cached = {
                "response": json.loads(content),
                "model": completion.model,
                "system_fingerprint": getattr(
                    completion, "system_fingerprint", None
                ),
            }
            dump_json(cache_path, cached)
        response_digest = canonical_digest(cached["response"])
        batch_entities, batch_relationships = normalize_spec_response(
            cached["response"],
            batch,
            model=cached.get("model", model),
            response_digest=response_digest,
            system_fingerprint=cached.get("system_fingerprint"),
        )
        entities.extend(batch_entities)
        relationships.extend(batch_relationships)
        response_digests.append(response_digest)
        fingerprints.append(cached.get("system_fingerprint"))
    response_digest = canonical_digest(response_digests)
    write_jsonl(graph / "spec_entities.jsonl", entities)
    write_jsonl(graph / "spec_relationships.jsonl", relationships)
    manifest["producers"]["spec"] = {
        "status": "passed",
        "model": model,
        "prompt_digest": prompt_digest,
        "input_digest": input_digest,
        "response_digest": response_digest,
        "batch_count": len(batches),
        "system_fingerprints": sorted({
            value for value in fingerprints if value
        }),
        "entity_count": len(entities),
        "relationship_count": len(relationships),
    }
    dump_json(graph / "manifest.json", manifest)
    return manifest["producers"]["spec"]


def spec_main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="eda-spec-produce")
    parser.add_argument("project")
    args = parser.parse_args(argv)
    try:
        result = produce_spec_graph(Path(args.project).resolve())
        print(json.dumps(result, indent=2, sort_keys=True, ensure_ascii=False))
        return 0
    except (FileNotFoundError, RuntimeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
